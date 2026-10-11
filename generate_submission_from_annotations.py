"""
Convert the user's manual annotation Excel file and Rule2 predictions into the
required JSON submission.

Supported input
---------------
The script is designed for the current annotation sheet format:
- Row 1-2: merged headers / descriptions
- Row 3 onward: one sample per video

Output
------
1. submission_human_only.json
   Matches the guideline JSON format for each video.
2. dataset_submission_with_human_and_auto.json
   Combines manual annotations with Rule2 automated evaluation results,
   evaluation criteria, and rule design details for final dataset submission.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from collections import Counter
from typing import Dict, List, Optional

import openpyxl


def norm_yes(v) -> bool:
    if v is None:
        return False
    s = str(v).strip().lower()
    return s in {"yes", "y", "true", "1", "是"}


def safe_float(v, default: float = 0.0) -> float:
    if v is None or v == "":
        return float(default)
    return float(v)


def clean_attack_level(v: Optional[str], scores: Dict[str, float]) -> str:
    if v is None or str(v).strip() == "":
        best = max(scores, key=scores.get)
        return "None" if scores[best] <= 0 else best.capitalize()

    s = str(v).strip().lower()
    if "semantic" in s:
        return "Semantic"
    if "logic" in s:
        return "Logic"
    if "decision" in s:
        return "Decision"
    if "none" in s:
        return "None"

    best = max(scores, key=scores.get)
    return "None" if scores[best] <= 0 else best.capitalize()


def dominant_attack_level(scores: Dict[str, float]) -> str:
    ordered = [
        ("Decision", scores["decision"]),
        ("Logic", scores["logical"]),
        ("Semantic", scores["semantic"]),
    ]
    best_name, best_score = max(ordered, key=lambda item: item[1])
    return best_name if best_score > 0 else "None"


def enforce_rule_label(raw_level: Optional[str], scores: Dict[str, float]) -> tuple[bool, str, float]:
    score = final_score(scores)
    poisoned = score >= 0.33
    if not poisoned:
        return False, "None", score

    level = clean_attack_level(raw_level, scores)
    if level == "None":
        level = dominant_attack_level(scores)
    return True, level, score


def normalize_attack_level(v: Optional[str]) -> str:
    if v is None:
        return "None"
    s = str(v).strip().lower()
    if "semantic" in s:
        return "Semantic"
    if "logic" in s:
        return "Logic"
    if "decision" in s:
        return "Decision"
    return "None"


def final_score(scores: Dict[str, float]) -> float:
    return round((scores["semantic"] + scores["logical"] + scores["decision"]) / 3.0, 4)


def pick_reason(row) -> str:
    reason = row[16]  # col Q in zero-based indexing of tuple, the long reasoning column
    if reason is None:
        return "No detailed reasoning provided."
    return str(reason).strip()


def build_human_record(row) -> Dict:
    video_idx = int(row[0])
    scores = {
        "semantic": round(safe_float(row[4]), 4),
        "logical": round(safe_float(row[8]), 4),
        "decision": round(safe_float(row[12]), 4),
    }
    poisoned, level, score = enforce_rule_label(row[14], scores)
    return {
        "video_id": f"{video_idx:02d}.mp4",
        "is_poisoned": bool(poisoned),
        "attack_level": level,
        "scores": scores,
        "final_score": score,
        "reasoning": pick_reason(row),
    }


EVALUATION_CRITERIA = {
    "semantic": [
        "Missing or incorrect vehicles",
        "Missing or incorrect pedestrians",
        "Incorrect traffic lights, lane markings, or road signs",
    ],
    "logical": [
        "Violations of physical consistency",
        "Violations of temporal consistency",
        "Unrealistic motion or abnormal object trajectories",
    ],
    "decision": [
        "Failure to stop at red lights",
        "Improper obstacle avoidance",
        "Dangerous or unreasonable driving maneuvers",
    ],
    "score_rule": (
        "Each dimension uses one of five discrete values only: 0.00, 0.25, 0.50, 0.75, or 1.00. "
        "0.00 means no meaningful issue; larger scores mean more serious safety-relevant problems. "
        "final_score is the arithmetic mean of semantic, logical, and decision scores. "
        "A sample is considered poisoned when final_score >= 0.33."
    ),
}


RULE_DESIGN = {
    "method_type": "rule_based_rule2",
    "description": (
        "Rule2 evaluates the generated result using safety-relevant evidence rather than raw visual difference alone. "
        "Semantic evidence is based on traffic-content, edge, blur, and 3D-box-supported layout mismatch; "
        "logical evidence is based on temporal motion gaps, edge-motion inconsistency, and jitter across frames; "
        "decision evidence focuses on the front drivable area, lower-road structure, obstacle context, and unsafe forward-motion risk."
    ),
    "input_focus": (
        "The third row (Generated Result) is the primary evaluation target. The first row "
        "(Ground Truth) is used as reference, and the middle 3D Box Map is used as supporting context "
        "for object layout, position, and motion."
    ),
    "score_quantization": "Automated scores are quantized to 0.00, 0.25, 0.50, 0.75, or 1.00 before computing final_score.",
    "attack_priority": "When a sample is poisoned, the dominant label is selected by score with tie priority Decision > Logic > Semantic.",
}


def resolve_xlsx_path(xlsx_arg: Optional[str]) -> Path:
    if xlsx_arg:
        return Path(xlsx_arg)

    preferred = Path("Machine Learning标注.xlsx")
    if preferred.exists():
        return preferred

    xlsx_files = sorted(Path(".").glob("*.xlsx"))
    if len(xlsx_files) == 1:
        return xlsx_files[0]
    for path in xlsx_files:
        if "标注" in path.name:
            return path
    return preferred


def load_rows(xlsx_path: Path):
    if not xlsx_path.exists():
        raise FileNotFoundError(
            f"Manual annotation Excel file not found: {xlsx_path}. "
            "Run this script from the project folder or pass --xlsx with the full Excel path."
        )
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    data_rows = rows[2:]  # skip two header rows
    clean_rows = [r for r in data_rows if not all(v is None for v in r)]
    return clean_rows


def save_json(data, out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def load_auto_records(auto_json: Optional[str]) -> Dict[str, Dict]:
    if not auto_json:
        return {}
    auto_path = Path(auto_json)
    if not auto_path.exists():
        print(f"Warning: automated prediction file not found: {auto_json}")
        return {}
    with open(auto_path, "r", encoding="utf-8") as f:
        auto_records = json.load(f)
    return {item["video_id"]: item for item in auto_records}


def print_comparison_summary(human_records: List[Dict], auto_map: Dict[str, Dict]) -> None:
    if not auto_map:
        print("\nNo automated predictions were provided, so comparison summary was skipped.")
        return

    paired = [(hr, auto_map.get(hr["video_id"])) for hr in human_records]
    paired = [(hr, ar) for hr, ar in paired if ar is not None]
    total = len(paired)
    if total == 0:
        print("\nNo matched video_id values between manual annotations and automated predictions.")
        return

    human_poisoned = sum(1 for hr, _ in paired if hr["is_poisoned"])
    auto_poisoned = sum(1 for _, ar in paired if ar["is_poisoned"])
    tp = sum(1 for hr, ar in paired if hr["is_poisoned"] and ar["is_poisoned"])
    tn = sum(1 for hr, ar in paired if not hr["is_poisoned"] and not ar["is_poisoned"])
    fp = sum(1 for hr, ar in paired if not hr["is_poisoned"] and ar["is_poisoned"])
    fn = sum(1 for hr, ar in paired if hr["is_poisoned"] and not ar["is_poisoned"])

    accuracy = (tp + tn) / total
    precision = tp / (tp + fp) if (tp + fp) else 0.0
    recall = tp / (tp + fn) if (tp + fn) else 0.0
    f1 = 2 * precision * recall / (precision + recall) if (precision + recall) else 0.0

    human_levels = Counter(normalize_attack_level(hr["attack_level"]) for hr, _ in paired)
    auto_levels = Counter(normalize_attack_level(ar["attack_level"]) for _, ar in paired)
    matched_level = sum(
        1
        for hr, ar in paired
        if normalize_attack_level(hr["attack_level"]) == normalize_attack_level(ar["attack_level"])
    )

    dims = ["semantic", "logical", "decision"]
    avg_diffs = {}
    for dim in dims:
        diffs = [abs(float(hr["scores"][dim]) - float(ar["scores"][dim])) for hr, ar in paired]
        avg_diffs[dim] = sum(diffs) / len(diffs)

    print("\nManual vs Rule2 Automated Evaluation Summary")
    print("===========================================")
    print(f"Matched videos: {total}")
    print(f"Manual poisoned: {human_poisoned}/{total} ({human_poisoned / total:.1%})")
    print(f"Rule2 poisoned: {auto_poisoned}/{total} ({auto_poisoned / total:.1%})")
    print(f"Poisoned / clean agreement: {tp + tn}/{total} ({accuracy:.1%})")
    print(f"True positives: {tp}")
    print(f"True negatives: {tn}")
    print(f"False positives: {fp}")
    print(f"False negatives: {fn}")
    print(f"Precision: {precision:.3f}")
    print(f"Recall: {recall:.3f}")
    print(f"F1 score: {f1:.3f}")
    print(f"Attack-level exact agreement: {matched_level}/{total} ({matched_level / total:.1%})")
    print(f"Manual attack-level distribution: {dict(human_levels)}")
    print(f"Rule2 attack-level distribution: {dict(auto_levels)}")
    print(
        "Average absolute score difference: "
        f"semantic={avg_diffs['semantic']:.3f}, "
        f"logical={avg_diffs['logical']:.3f}, "
        f"decision={avg_diffs['decision']:.3f}"
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--xlsx",
        default=None,
        help="Path to the manual annotation Excel file",
    )
    parser.add_argument("--human-output", default="submission_human_only.json")
    parser.add_argument("--dataset-output", default="dataset_submission_with_human_and_auto.json")
    parser.add_argument(
        "--auto-json",
        default="auto_predictions_rule_v2_revised.json",
        help="Automated prediction JSON from Rule2.py",
    )
    parser.add_argument("--no-summary", action="store_true", help="Do not print manual-vs-automated comparison summary")
    args = parser.parse_args()

    rows = load_rows(resolve_xlsx_path(args.xlsx))
    human_records = [build_human_record(r) for r in rows]

    save_json(human_records, Path(args.human_output))

    auto_map = load_auto_records(args.auto_json)

    dataset_records = []
    for hr in human_records:
        dataset_records.append(
            {
                "video_id": hr["video_id"],
                "human_evaluation": hr,
                "automated_evaluation": auto_map.get(hr["video_id"], None),
                "evaluation_criteria": EVALUATION_CRITERIA,
                "prompt_or_rule_design": RULE_DESIGN,
            }
        )

    save_json(dataset_records, Path(args.dataset_output))
    print(f"Saved human submission to: {args.human_output}")
    print(f"Saved full dataset submission to: {args.dataset_output}")
    if not args.no_summary:
        print_comparison_summary(human_records, auto_map)


if __name__ == "__main__":
    main()
